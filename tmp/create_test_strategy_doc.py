from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/carolinaecheverri/Documents/APIS/logistics-management")
OUT = ROOT / "output" / "docs"
TMP = ROOT / "tmp" / "docs"
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / "Estrategia_y_escenarios_pruebas_colas_mensajes.docx"
DIAGRAM_PATH = TMP / "arquitectura_colas.png"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F3F6FA"
GRAY = "D9D9D9"
TEXT = "202020"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, lines, fill, outline=NAVY):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline="#" + outline, width=4)
    title_font = font(30, True)
    body_font = font(22)
    draw.text(((x1 + x2) / 2, y1 + 24), title, anchor="ma", fill="#000000", font=title_font)
    y = y1 + 72
    for line in lines:
        draw.text(((x1 + x2) / 2, y), line, anchor="ma", fill="#202020", font=body_font)
        y += 32


def arrow(draw, start, end, label=None, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        steps = 14
        for i in range(0, steps, 2):
            a = i / steps
            b = min((i + 1) / steps, 1)
            draw.line((x1 + (x2 - x1) * a, y1 + (y2 - y1) * a,
                       x1 + (x2 - x1) * b, y1 + (y2 - y1) * b), fill="#17365D", width=5)
    else:
        draw.line((x1, y1, x2, y2), fill="#17365D", width=5)
    import math
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 22
    for delta in (2.55, -2.55):
        draw.line((x2, y2, x2 + length * math.cos(angle + delta),
                   y2 + length * math.sin(angle + delta)), fill="#17365D", width=5)
    if label:
        draw.text(((x1 + x2) / 2, (y1 + y2) / 2 - 17), label, anchor="ms",
                  fill="#000000", font=font(20, True))


def build_diagram():
    image = Image.new("RGB", (1800, 930), "white")
    draw = ImageDraw.Draw(image)
    rounded_box(draw, (60, 160, 320, 390), "Cliente", ["Swagger", "o curl"], "#F3F6FA")
    rounded_box(draw, (400, 110, 760, 440), "shipment-service", ["API REST", "reglas de negocio", "productor AMQP"], "#DCE6F1")
    rounded_box(draw, (870, 110, 1220, 440), "RabbitMQ", ["shipment.exchange", "cola principal", "routing key"], "#E2F0D9")
    rounded_box(draw, (1330, 110, 1730, 440), "notification-service", ["consumidor", "3 intentos", "envio de correo"], "#FFF2CC")
    rounded_box(draw, (440, 610, 720, 820), "PostgreSQL", ["envios", "idempotencia"], "#F3F6FA")
    rounded_box(draw, (900, 610, 1190, 820), "DLQ", ["mensajes", "no procesados"], "#FCE4D6")
    rounded_box(draw, (1370, 610, 1690, 820), "SMTP", ["Mailpit", "o Gmail"], "#F3F6FA")
    arrow(draw, (320, 275), (400, 275), "HTTP")
    arrow(draw, (760, 275), (870, 275), "evento JSON")
    arrow(draw, (1220, 275), (1330, 275), "entrega")
    arrow(draw, (580, 440), (580, 610), "persistencia")
    arrow(draw, (1045, 440), (1045, 610), "rechazo final", dashed=True)
    arrow(draw, (1530, 440), (1530, 610), "SMTP")
    draw.text((900, 40), "Flujo de integración asíncrona", anchor="ma", fill="#000000", font=font(38, True))
    image.save(DIAGRAM_PATH, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths=None, header=True, font_size=9.2):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and c_idx < len(widths):
                cell.width = Inches(widths[c_idx])
            if header and r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor(255, 255, 255) if header and r_idx == 0 else RGBColor.from_string(TEXT)
                    if header and r_idx == 0:
                        run.bold = True
    if header:
        repeat_header(table.rows[0])


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    style_table(table, widths, True, font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_detail_table(doc, data):
    table = doc.add_table(rows=0, cols=2)
    for i, (label, value) in enumerate(data):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], BLUE)
        if i % 2 == 1:
            set_cell_shading(cells[1], PALE)
        for c_idx, cell in enumerate(cells):
            set_cell_border(cell)
            set_cell_margins(cell, top=90, bottom=90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(1.6 if c_idx == 0 else 5.25)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor.from_string(TEXT)
                    if c_idx == 0:
                        run.bold = True
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


NUMBER_COUNTER = 0


def reset_numbering():
    global NUMBER_COUNTER
    NUMBER_COUNTER = 0


def add_number(doc, text):
    global NUMBER_COUNTER
    NUMBER_COUNTER += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.add_run(f"{NUMBER_COUNTER}.  {text}")
    return p


def add_code(doc, text):
    for line in text.strip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line)
        run.font.name = "Courier New"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
        run.font.size = Pt(9)


def add_page_number(paragraph):
    run = paragraph.add_run(" | Pagina ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.85)
    section.footer_distance = Inches(0.3)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = ""


def add_case(doc, title, data):
    heading = doc.add_heading(title, level=2)
    heading.paragraph_format.page_break_before = True
    add_detail_table(doc, data)


def build_document():
    build_diagram()
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Estrategia y escenarios de prueba para arquitectura basada en colas de mensajes")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Laboratorio Logistics Management")
    r.bold = True
    r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Evaluacion del tema 3 de la guia de exposicion y laboratorio\nFecha de revision 3 de octubre de 2026")
    doc.add_paragraph()
    doc.add_heading("Conclusion del diagnostico", level=1)
    doc.add_paragraph(
        "El proyecto cumple el nucleo tecnico solicitado para el tema 3. Tiene un productor, "
        "RabbitMQ como intermediario, una cola principal, un consumidor, confirmacion del consumo "
        "al finalizar correctamente, tres intentos ante errores y una cola de mensajes no procesados. "
        "La ejecucion automatizada realizada para esta revision termino con 10 pruebas, 0 fallos y "
        "0 pruebas omitidas."
    )
    doc.add_paragraph(
        "La entrega academica aun no esta completa por si sola. Antes de exponer se deben guardar "
        "capturas o reportes de la ejecucion, preparar la presentacion y mostrar de forma visible la "
        "confirmacion del mensaje. Tambien conviene automatizar un flujo completo que llegue hasta un "
        "servidor SMTP de laboratorio."
    )
    doc.add_paragraph().add_run("Dictamen: cumple tecnicamente; entrega documental y evidencias en preparacion.").bold = True
    doc.add_page_break()

    doc.add_heading("Arquitectura evaluada", level=1)
    doc.add_paragraph(
        "El laboratorio agrupa dos aplicaciones Spring Boot independientes. shipment-service recibe "
        "solicitudes HTTP, guarda el envio en PostgreSQL y publica un evento. RabbitMQ enruta el mensaje "
        "hacia shipment.notification.queue. notification-service consume el evento y solicita el envio "
        "de correo. Si el procesamiento falla tres veces, el mensaje se rechaza sin reencolar y pasa a "
        "shipment.notification.dlq."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(DIAGRAM_PATH), width=Inches(6.75))
    p = doc.add_paragraph("Figura 1  Arquitectura y flujo de mensajes del laboratorio")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()
    doc.add_heading("Verificacion del tema 3", level=1)
    compliance_rows = [
        ("Productor", "RabbitShipmentEventPublisher publica con RabbitTemplate", "Cumple", "Conservar evidencia del evento publicado"),
        ("Intermediario", "RabbitMQ declarado en docker-compose.yml", "Cumple", "Mostrar el contenedor saludable"),
        ("Cola", "shipment.notification.queue es durable", "Cumple", "Mostrarla en RabbitMQ Management"),
        ("Consumidor", "ShipmentEventConsumer usa RabbitListener", "Cumple", "Relacionar eventId entre productor y consumidor"),
        ("Envio y recepcion", "Evento JSON publicado y convertido a ShipmentEvent", "Cumple", "Capturar respuesta y logs"),
        ("Confirmacion", "El listener termina normalmente y el contenedor confirma el consumo", "Cumple con evidencia pendiente", "Mostrar Ready y Unacked en cero tras el exito"),
        ("Reintentos", "Tres intentos totales con esperas de 1 y 2 segundos", "Cumple", "Capturar los logs intento 1/3 a 3/3"),
        ("No procesables", "RejectAndDontRequeueRecoverer y DLQ", "Cumple", "Capturar incremento de la DLQ"),
        ("Automatizacion", "JUnit, Mockito y Testcontainers", "Cumple", "Adjuntar reportes HTML"),
    ]
    add_table(doc, ["Requisito", "Evidencia en el proyecto", "Estado", "Accion para la exposicion"],
              compliance_rows, [1.15, 2.45, 1.25, 2.0], 8.4)

    doc.add_heading("Aspectos que faltan para completar la entrega", level=1)
    add_bullet(doc, "Presentacion de aproximadamente 20 minutos. No existe una presentacion dentro del repositorio.")
    add_bullet(doc, "Evidencias exportadas. Los reportes se generan en build, pero deben guardarse capturas o copias para la entrega.")
    add_bullet(doc, "Prueba automatizada de extremo a extremo hasta Mailpit. Las pruebas actuales validan productor y cola, o consumidor con dependencias simuladas, pero no todo el recorrido en una sola prueba.")
    add_bullet(doc, "Evidencia directa del ACK. El comportamiento existe por el modo del contenedor, pero debe mostrarse con la cola vacia y sin mensajes Unacked.")
    add_bullet(doc, "Video corto de respaldo, recomendado por la guia para problemas tecnicos durante la demostracion.")
    add_bullet(doc, "Como mejora adicional, habilitar publisher confirms si se desea demostrar tambien la confirmacion entre el productor y RabbitMQ.")
    doc.add_heading("Estrategia de pruebas", level=1)
    doc.add_heading("Objetivo", level=2)
    doc.add_paragraph(
        "Validar que los componentes intercambian eventos de forma asincrona, que un mensaje correcto "
        "se procesa una sola vez y que los errores se controlan mediante validacion, reintentos y DLQ "
        "sin detener el consumidor ni producir envios de correo repetidos."
    )
    doc.add_heading("Componentes incluidos", level=2)
    add_bullet(doc, "shipment-service como API, persistencia y productor de mensajes.")
    add_bullet(doc, "RabbitMQ con exchange, routing key, cola principal, dead letter exchange y DLQ.")
    add_bullet(doc, "notification-service como consumidor y generador de notificaciones.")
    add_bullet(doc, "PostgreSQL para persistencia e idempotencia de solicitudes.")
    add_bullet(doc, "Mailpit como SMTP simulado recomendado, o Gmail con una cuenta propia autorizada.")

    doc.add_heading("Puntos de integracion", level=2)
    add_bullet(doc, "HTTP entre el cliente y shipment-service.")
    add_bullet(doc, "JPA entre shipment-service y PostgreSQL.")
    add_bullet(doc, "AMQP entre shipment-service, RabbitMQ y notification-service.")
    add_bullet(doc, "SMTP entre notification-service y el servidor de correo.")

    doc.add_heading("Riesgos identificados", level=2)
    risk_rows = [
        ("R1", "El productor no publica o el mensaje no se enruta", "Alta", "Prueba de publicacion y lectura desde la cola"),
        ("R2", "El consumidor procesa dos veces el mismo eventId", "Alta", "Prueba de duplicado e idempotencia"),
        ("R3", "Un error genera reintentos infinitos", "Alta", "Limite de tres intentos y DLQ"),
        ("R4", "Un mensaje invalido detiene el consumidor", "Alta", "Validacion, rechazo y continuidad"),
        ("R5", "El consumidor no esta disponible", "Media", "Acumulacion en Ready y consumo al reiniciar"),
        ("R6", "El SMTP no responde", "Media", "Error controlado, reintento y DLQ"),
        ("R7", "Una solicitud repetida crea dos envios", "Alta", "Idempotency-Key y restriccion unica"),
    ]
    add_table(doc, ["ID", "Riesgo", "Impacto", "Tratamiento mediante prueba"], risk_rows,
              [0.45, 2.65, 0.75, 3.05], 8.8)

    doc.add_heading("Alcance y exclusiones", level=2)
    doc.add_paragraph(
        "El alcance comprende la ejecucion local con Docker Compose, la API de envios, PostgreSQL, "
        "RabbitMQ, el consumidor y un SMTP de laboratorio. No incluye carga, alta disponibilidad de "
        "RabbitMQ, despliegue en nube, seguridad ofensiva ni pruebas sobre servicios de terceros. Gmail "
        "solo se usa con una cuenta propia y datos ficticios."
    )

    doc.add_heading("Ambiente y herramientas", level=2)
    environment_rows = [
        ("Java", "21", "Ejecucion de los microservicios"),
        ("Spring Boot", "4.1.1", "API, AMQP, JPA y correo"),
        ("RabbitMQ", "4 management alpine", "Broker y consola de administracion"),
        ("PostgreSQL", "17 alpine", "Persistencia"),
        ("Docker Compose", "Local", "Orquestacion del laboratorio"),
        ("JUnit y Mockito", "Gradle", "Pruebas automatizadas"),
        ("Testcontainers", "PostgreSQL y RabbitMQ", "Dependencias efimeras de integracion"),
        ("Swagger", "localhost 8080", "Ejecucion manual de la API"),
        ("Mailpit", "localhost 8025", "Correo simulado y evidencia visual"),
    ]
    add_table(doc, ["Herramienta", "Version o ubicacion", "Uso"], environment_rows, [1.5, 2.0, 3.4], 9)

    doc.add_heading("Datos de prueba", level=2)
    add_bullet(doc, "Correos ficticios como prueba@example.com para escenarios normales con Mailpit.")
    add_bullet(doc, "fail@email.com para activar el fallo controlado y demostrar los tres intentos.")
    add_bullet(doc, "Claves de idempotencia unicas por caso, por ejemplo lab-cp01-001.")
    add_bullet(doc, "Nombres, origenes, destinos y descripciones ficticios, sin informacion personal.")

    doc.add_heading("Dependencias reales y simuladas", level=2)
    doc.add_paragraph(
        "RabbitMQ y PostgreSQL se ejecutan como dependencias reales en contenedores. Testcontainers "
        "crea instancias efimeras para las pruebas de integracion. Mockito simula los puertos de "
        "persistencia, publicacion y correo en pruebas de aplicacion. Mailpit simula el servicio SMTP."
    )

    doc.add_heading("Criterios de aprobacion y fallo", level=2)
    add_bullet(doc, "Una prueba se aprueba cuando termina sin excepciones no esperadas y todas sus aserciones coinciden con el resultado esperado.")
    add_bullet(doc, "Las pruebas de mensajeria deben conservar el mismo eventId y shipmentId en productor y consumidor.")
    add_bullet(doc, "El escenario exitoso debe retirar el mensaje de la cola y registrar la confirmacion del consumo.")
    add_bullet(doc, "El escenario fallido debe ejecutar exactamente tres intentos y terminar en la DLQ, sin requeue infinito.")
    add_bullet(doc, "Una solicitud o evento duplicado no debe crear un segundo envio ni enviar un segundo correo.")
    add_bullet(doc, "Cualquier fallo de asercion, mensaje perdido, duplicacion o cola en estado inesperado hace fallar el caso.")
    doc.add_heading("Matriz de casos de prueba", level=1)
    matrix_rows = [
        ("CP-01", "Publicacion correcta en RabbitMQ", "Integracion", "Si", "R1", "Aprobado"),
        ("CP-02", "Solicitud con correo invalido", "Integracion", "Si", "R4", "Aprobado"),
        ("CP-03", "Clave idempotente con datos diferentes", "Sistema", "Parcial", "R7", "Parcial"),
        ("CP-04", "Tres intentos y envio a DLQ", "Integracion", "Si", "R3 R6", "Aprobado"),
        ("CP-05", "Evento duplicado sin segundo correo", "Integracion", "Si", "R2", "Aprobado"),
        ("CP-06", "Consumidor detenido y recuperacion", "Sistema", "No", "R5", "Pendiente"),
        ("CP-07", "Flujo completo hasta Mailpit", "Sistema", "No", "R1 R6", "Pendiente"),
    ]
    add_table(doc, ["ID", "Escenario", "Tipo", "Automatizada", "Riesgo", "Estado"], matrix_rows,
              [0.55, 2.6, 1.0, 1.0, 0.75, 1.0], 8.5)
    doc.add_paragraph(
        "Los casos CP-01, CP-02 y CP-04 satisfacen el minimo de automatizacion solicitado por la guia: "
        "un escenario exitoso, uno de validacion y uno de falla tecnica."
    )

    add_case(doc, "Escenario 01 - Publicacion correcta en RabbitMQ", [
        ("Identificador", "Caso CP-01"),
        ("Nombre", "Persistencia y publicacion de un envio creado"),
        ("Tipo de prueba", "Integracion"),
        ("Objetivo", "Comprobar que shipment-service guarda el envio y publica un unico evento SHIPMENT_CREATED."),
        ("Componentes", "CreateShipmentUseCase, PostgreSQL, RabbitShipmentEventPublisher y RabbitMQ."),
        ("Precondiciones", "Docker disponible. Contenedores efimeros de PostgreSQL y RabbitMQ iniciados por Testcontainers."),
        ("Datos de entrada", "Idempotency-Key integration-create-1 y datos ficticios de un envio."),
        ("Procedimiento", "Ejecutar la prueba de integracion de infraestructura. Crear el envio dos veces con la misma clave. Consultar la base y leer la cola."),
        ("Resultado esperado", "Una fila persistida, un tracking SHP, un solo evento en la cola y ningun segundo mensaje por la repeticion."),
        ("Resultado obtenido", "La prueba termino aprobada el 3 de octubre de 2026."),
        ("Evidencia", "shipment-service/build/reports/tests/test/index.html y ShipmentInfrastructureIntegrationTest."),
        ("Estado", "Aprobado"),
    ])

    add_case(doc, "Escenario 02 - Solicitud con correo invalido", [
        ("Identificador", "Caso CP-02"),
        ("Nombre", "Rechazo de una solicitud HTTP con correo invalido"),
        ("Tipo de prueba", "Integracion"),
        ("Objetivo", "Validar que el controller responde 400 y no invoca el caso de uso cuando el correo no tiene formato valido."),
        ("Componentes", "ShipmentController, validacion Jakarta y GlobalExceptionHandler."),
        ("Precondiciones", "Aplicacion de prueba construida con MockMvc."),
        ("Datos de entrada", "customerEmail invalid y los demas campos obligatorios."),
        ("Procedimiento", "Ejecutar la prueba automatizada de correo invalido en ShipmentControllerValidationTest."),
        ("Resultado esperado", "HTTP 400, cuerpo de error y cero interacciones con CreateShipmentPort."),
        ("Resultado obtenido", "La prueba termino aprobada el 3 de octubre de 2026."),
        ("Evidencia", "shipment-service/build/reports/tests/test/index.html y ShipmentControllerValidationTest."),
        ("Estado", "Aprobado"),
    ])

    add_case(doc, "Escenario 03 - Conflicto de idempotencia", [
        ("Identificador", "Caso CP-03"),
        ("Nombre", "Reutilizacion de una clave con datos diferentes"),
        ("Tipo de prueba", "Sistema"),
        ("Objetivo", "Evitar que una misma clave idempotente cree dos envios o acepte datos incompatibles."),
        ("Componentes", "API de shipment-service, caso de uso de creacion y PostgreSQL."),
        ("Precondiciones", "Existe un envio creado con la clave lab-cp03-001."),
        ("Datos de entrada", "Segundo POST con la misma clave y un customerName diferente."),
        ("Procedimiento", "Crear el primer envio. Repetir el POST cambiando un campo. Consultar la lista de envios."),
        ("Resultado esperado", "HTTP 409 y una sola fila asociada a la clave. No se publica otro evento."),
        ("Resultado obtenido", "La regla fue aprobada en ShipmentUseCasesTest. Falta guardar evidencia de la ejecucion HTTP completa."),
        ("Evidencia", "Prueba de conflicto de idempotencia en ShipmentUseCasesTest y captura pendiente de Swagger."),
        ("Estado", "Parcial"),
    ])

    add_case(doc, "Prueba de reintentos y DLQ (CP-04)", [
        ("Codigo del caso", "CP-04"),
        ("Nombre", "Mensaje que no puede procesarse despues de tres intentos"),
        ("Tipo de prueba", "Integracion"),
        ("Objetivo", "Comprobar que un error controlado no produce un ciclo infinito y que el mensaje termina en la DLQ."),
        ("Componentes", "RabbitMQ, listener, retry interceptor, ProcessShipmentNotificationUseCase y DLQ."),
        ("Precondiciones", "RabbitMQ efimero iniciado por Testcontainers."),
        ("Datos de entrada", "Evento valido con customerEmail fail@email.com."),
        ("Procedimiento", "Ejecutar la prueba de integracion de retry y DLQ en NotificationRabbitIntegrationTest."),
        ("Resultado esperado", "Intentos 1/3, 2/3 y 3/3; rechazo sin requeue; mensaje disponible en shipment.notification.dlq."),
        ("Resultado obtenido", "La prueba termino aprobada el 3 de octubre de 2026 y el mensaje fue leido desde la DLQ."),
        ("Evidencia", "notification-service/build/reports/tests/test/index.html, NotificationRabbitIntegrationTest y logs del consumidor."),
        ("Estado", "Aprobado"),
    ])

    add_case(doc, "Escenario 05 - Evento duplicado", [
        ("Identificador", "Caso CP-05"),
        ("Nombre", "Procesamiento idempotente por eventId"),
        ("Tipo de prueba", "Integracion con dependencia SMTP simulada"),
        ("Objetivo", "Comprobar que dos entregas con el mismo eventId generan un solo correo."),
        ("Componentes", "ProcessShipmentNotificationUseCase y NotificationSenderPort simulado."),
        ("Precondiciones", "Instancia nueva del caso de uso y sender creado con Mockito."),
        ("Datos de entrada", "El mismo ShipmentEvent se procesa dos veces."),
        ("Procedimiento", "Ejecutar la prueba automatizada de evento duplicado en ProcessShipmentNotificationUseCaseTest."),
        ("Resultado esperado", "NotificationSenderPort recibe una sola invocacion y la segunda entrega se ignora."),
        ("Resultado obtenido", "La prueba termino aprobada el 3 de octubre de 2026."),
        ("Evidencia", "notification-service/build/reports/tests/test/index.html y ProcessShipmentNotificationUseCaseTest."),
        ("Estado", "Aprobado"),
    ])

    add_case(doc, "Escenario 06 - Consumidor detenido y recuperacion", [
        ("Identificador", "Caso CP-06"),
        ("Nombre", "Acumulacion y consumo posterior del mensaje"),
        ("Tipo de prueba", "Sistema"),
        ("Objetivo", "Demostrar que RabbitMQ conserva el mensaje mientras notification-service no esta disponible."),
        ("Componentes", "shipment-service, RabbitMQ y notification-service."),
        ("Precondiciones", "Docker Compose activo y cola principal vacia."),
        ("Datos de entrada", "POST valido con una nueva Idempotency-Key y correo ficticio para Mailpit."),
        ("Procedimiento", "Detener notification-service. Crear el envio. Verificar un mensaje Ready. Iniciar notification-service y comprobar que Ready vuelve a cero."),
        ("Resultado esperado", "El productor responde 201, el mensaje permanece en la cola y se procesa cuando vuelve el consumidor."),
        ("Resultado obtenido", "Procedimiento documentado en README. Ejecucion y captura pendientes para la entrega."),
        ("Evidencia", "Capturas pendientes de docker compose ps, RabbitMQ Management y logs correlacionados."),
        ("Estado", "Pendiente"),
    ])

    add_case(doc, "Recorrido completo hacia Mailpit, caso CP-07", [
        ("Clave unica de prueba", "Identificador CP-07"),
        ("Nombre", "Creacion de envio y correo visible en SMTP de laboratorio"),
        ("Tipo de prueba", "Sistema"),
        ("Objetivo", "Demostrar el recorrido desde la solicitud HTTP hasta la notificacion sin usar datos reales."),
        ("Componentes", "Cliente, shipment-service, PostgreSQL, RabbitMQ, notification-service y Mailpit."),
        ("Precondiciones", "Compose configurado con MAIL_HOST=mailpit, MAIL_PORT=1025, autenticacion y STARTTLS desactivados."),
        ("Datos de entrada", "Envio valido con destinatario prueba@example.com y clave lab-cp07-001."),
        ("Procedimiento", "Crear el envio en Swagger. Seguir eventId y shipmentId en los logs. Abrir localhost:8025 y revisar el mensaje."),
        ("Resultado esperado", "HTTP 201, evento recibido, ACK, cola vacia y un correo con tracking, origen y destino."),
        ("Resultado obtenido", "Pendiente de ejecutar con Mailpit y guardar capturas."),
        ("Evidencia", "Capturas pendientes de Swagger, logs, RabbitMQ Management y bandeja de Mailpit."),
        ("Estado", "Pendiente"),
    ])
    doc.add_page_break()

    doc.add_heading("Ejecucion automatizada", level=1)
    doc.add_paragraph(
        "La ejecucion del 3 de octubre de 2026 uso el socket de Colima para que Testcontainers pudiera "
        "crear PostgreSQL y RabbitMQ efimeros. Los dos proyectos se ejecutaron por separado con su "
        "propio wrapper Gradle."
    )
    result_rows = [
        ("shipment-service", "6", "0", "0", "Aprobado"),
        ("notification-service", "4", "0", "0", "Aprobado"),
        ("Total", "10", "0", "0", "Aprobado"),
    ]
    add_table(doc, ["Proyecto", "Pruebas", "Fallos", "Omitidas", "Resultado"], result_rows,
              [2.2, 1.0, 1.0, 1.0, 1.35], 9.2)
    doc.add_heading("Comandos reproducibles", level=2)
    add_code(doc, """
cd shipment-service
DOCKER_HOST=unix:///Users/carolinaecheverri/.colima/default/docker.sock \\
TESTCONTAINERS_DOCKER_SOCKET_OVERRIDE=/var/run/docker.sock \\
./gradlew clean test

cd ../notification-service
DOCKER_HOST=unix:///Users/carolinaecheverri/.colima/default/docker.sock \\
TESTCONTAINERS_DOCKER_SOCKET_OVERRIDE=/var/run/docker.sock \\
./gradlew clean test
""")
    doc.add_paragraph(
        "Si Docker Desktop esta activo, normalmente basta con ejecutar ./gradlew clean test dentro "
        "de cada microservicio."
    )

    doc.add_heading("Evidencias que deben guardarse", level=1)
    evidence_rows = [
        ("E-01", "docker compose ps con todos los servicios activos", "Pendiente de captura"),
        ("E-02", "Solicitud y respuesta 201 desde Swagger", "Pendiente de captura"),
        ("E-03", "Mismo eventId en productor y consumidor", "Disponible en logs"),
        ("E-04", "Cola principal sin Ready ni Unacked despues del exito", "Pendiente de captura"),
        ("E-05", "Logs de intento 1/3, 2/3 y 3/3", "Disponible en logs"),
        ("E-06", "Mensaje visible en shipment.notification.dlq", "Pendiente de captura"),
        ("E-07", "Reportes HTML con 10 pruebas y 0 fallos", "Disponible en build/reports"),
        ("E-08", "Correo visible en Mailpit", "Pendiente de captura"),
    ]
    add_table(doc, ["ID", "Evidencia", "Situacion"], evidence_rows, [0.65, 4.4, 1.8], 9)
    doc.add_paragraph(
        "Las capturas deben usar datos ficticios y mostrar fecha, nombre de la cola, identificadores y "
        "resultado. Para conservar los reportes, copiar las carpetas build/reports/tests/test a una "
        "carpeta evidencias antes de ejecutar clean."
    )

    doc.add_heading("Guion de demostracion", level=1)
    reset_numbering()
    add_number(doc, "Mostrar la arquitectura y explicar productor, broker, cola, consumidor y DLQ en 3 minutos.")
    add_number(doc, "Presentar riesgos, estrategia y matriz de casos en 5 minutos.")
    add_number(doc, "Ejecutar CP-07 como escenario exitoso y mostrar el ACK mediante la cola vacia.")
    add_number(doc, "Ejecutar CP-04 con fail@email.com y seguir los tres intentos hasta la DLQ.")
    add_number(doc, "Abrir los reportes automatizados y explicar resultados y limitaciones en 3 minutos.")
    add_number(doc, "Reservar 1 minuto para preguntas, de acuerdo con la guia.")

    doc.add_heading("Instalacion y ejecucion del laboratorio", level=1)
    reset_numbering()
    add_number(doc, "Instalar Docker o Colima y verificar que el daemon este activo.")
    add_number(doc, "Para la exposicion, configurar Mailpit como SMTP y usar destinatarios ficticios.")
    add_number(doc, "Desde la raiz del repositorio ejecutar docker compose up -d --build.")
    add_number(doc, "Abrir Swagger en http://localhost:8080/swagger-ui/index.html.")
    add_number(doc, "Abrir RabbitMQ Management en http://localhost:15672 con guest y guest.")
    add_number(doc, "Abrir Mailpit en http://localhost:8025.")
    add_number(doc, "Seguir la ejecucion con docker compose logs -f shipment-service notification-service.")
    add_number(doc, "Al terminar ejecutar docker compose down. No agregar -v si se desea conservar PostgreSQL.")

    doc.add_heading("Analisis de resultados", level=1)
    doc.add_paragraph(
        "Los resultados automatizados demuestran que el productor persiste y publica, que la API "
        "rechaza datos invalidos y que un mensaje con fallo controlado llega a la DLQ. La prueba de "
        "duplicados demuestra que el puerto de correo se invoca una sola vez por eventId durante la "
        "vida de la instancia."
    )
    doc.add_paragraph(
        "La principal limitacion es que la deduplicacion del consumidor reside en memoria. Un reinicio "
        "borra los eventId procesados y varias replicas no compartirian el registro. Esto no impide "
        "demostrar el tema 3, pero debe explicarse como riesgo y posible mejora mediante una tabla de "
        "eventos procesados con restriccion unica."
    )
    doc.add_paragraph(
        "Otra limitacion es que el test exitoso del productor y el test del consumidor se ejecutan por "
        "separado. Una prueba futura debe levantar RabbitMQ y Mailpit, crear el envio por HTTP y esperar "
        "hasta encontrar el correo. Esa prueba daria evidencia automatizada de todo el flujo de sistema."
    )
    doc.add_paragraph(
        "El envio actual usa RabbitTemplate sin publisher confirms. La confirmacion exigida por el tema "
        "se puede demostrar como ACK del consumidor. Si el docente pide confirmacion del broker al "
        "productor, se deben habilitar publisher confirms y probar los callbacks de ack y nack."
    )

    doc.add_heading("Conclusiones", level=1)
    add_bullet(doc, "La arquitectura implementa productor, cola, consumidor, reintentos y DLQ, por lo que satisface el objetivo tecnico central del tema 3.")
    add_bullet(doc, "Las 10 pruebas ejecutadas aprobaron y cubren los tres escenarios automatizados obligatorios." )
    add_bullet(doc, "La calificacion dependera tambien de evidencias visibles, explicacion de los resultados y claridad de la demostracion." )
    add_bullet(doc, "Mailpit es la opcion recomendada para el laboratorio porque evita datos personales y dependencias externas." )

    doc.add_heading("Lista de entregables", level=1)
    deliverable_rows = [
        ("Presentacion", "Pendiente", "Crear diapositivas con el guion de 20 minutos"),
        ("Diagrama de arquitectura", "Incluido", "Figura 1 de este documento"),
        ("Estrategia de pruebas", "Incluida", "Seccion Estrategia de pruebas"),
        ("Matriz con minimo cinco casos", "Incluida", "Siete casos definidos"),
        ("Codigo fuente", "Disponible", "Dos microservicios independientes"),
        ("Scripts y configuracion", "Disponible", "docker-compose.yml, Dockerfiles y Gradle"),
        ("Instrucciones", "Disponibles", "README y este documento"),
        ("Evidencias ejecutadas", "Parcial", "Guardar capturas y copiar reportes"),
        ("Analisis y conclusiones", "Incluidos", "Secciones finales"),
        ("Referencias", "Incluidas", "Seccion Referencias"),
    ]
    add_table(doc, ["Entregable", "Estado", "Ubicacion o accion"], deliverable_rows, [2.4, 1.2, 3.25], 9)

    doc.add_heading("Referencias", level=1)
    refs = [
        "Jaramillo Patiño, Carlos Andres. Guia para la exposicion y laboratorio Estrategias de pruebas de integracion y de sistema. 2026.",
        "RabbitMQ. Consumer Acknowledgements and Publisher Confirms. https://www.rabbitmq.com/docs/confirms",
        "RabbitMQ. Dead Letter Exchanges. https://www.rabbitmq.com/docs/dlx",
        "Spring AMQP. Resilience Recovering from Errors and Broker Failures. https://docs.spring.io/spring-amqp/reference/amqp/resilience-recovering-from-errors-and-broker-failures.html",
        "Testcontainers for Java. RabbitMQ module. https://java.testcontainers.org/modules/rabbitmq/",
        "Repositorio Logistics Management. README, codigo fuente, configuraciones y reportes de pruebas. Revision del 3 de octubre de 2026.",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
